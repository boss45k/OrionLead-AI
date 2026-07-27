import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.oxml.ns import qn

doc = Document(r'C:\AI-Lead-Collection-System\report\OrionLead_AI_Report.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print("=" * 60)

page = 1
for i, para in enumerate(doc.paragraphs[:80]):
    text = para.text.strip()
    style = para.style.name

    # Detect page breaks
    for elem in para._element.iter():
        if elem.tag == qn('w:br'):
            btype = elem.get(qn('w:type'), '')
            if btype == 'page':
                page += 1
                print(f"\n{'='*60} PAGE BREAK → Page {page} {'='*60}\n")
        if elem.tag == qn('w:sectPr'):
            page += 1
            print(f"\n{'='*60} SECTION BREAK → Page {page} {'='*60}\n")

    # Check for images
    has_image = any(e.tag == qn('w:drawing') or e.tag == qn('w:pict')
                    for e in para._element.iter())

    label = f"[{i:3d}] style='{style}'"
    if has_image:
        label += " [IMAGE]"
    if text:
        label += f" | '{text[:80]}'"
    else:
        label += " | (empty)"

    print(label)

    if page > 3:
        print("\n... stopping at page 3")
        break
