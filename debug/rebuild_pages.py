"""
Rebuild first two pages of OrionLead_AI_Report.docx to match reference layout:
  Page 1: IUL logo → project title/author/details
  Page 2: IUL logo → defense committee block → signature lines
"""
import shutil, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC    = r'C:\AI-Lead-Collection-System\report\OrionLead_AI_Report.docx'
BACKUP = r'C:\AI-Lead-Collection-System\report\OrionLead_AI_Report_BACKUP2.docx'

shutil.copy(SRC, BACKUP)
print("Backup saved.")

doc = Document(SRC)

# ── Helpers ───────────────────────────────────────────────────────────────────────

def clear_runs(para):
    """Remove all content from paragraph while preserving pPr (paragraph properties)."""
    p = para._element
    pPr = p.find(qn('w:pPr'))
    for child in list(p):
        p.remove(child)
    if pPr is not None:
        p.insert(0, pPr)

def delete_para(para):
    para._element.getparent().remove(para._element)

def fmt_para(para, align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=0):
    para.alignment = align
    para.paragraph_format.space_before = Pt(sb)
    para.paragraph_format.space_after  = Pt(sa)
    para.paragraph_format.line_spacing = None

def set_text(para, text, bold=False, size=12,
             align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=6):
    clear_runs(para)
    fmt_para(para, align=align, sb=sb, sa=sa)
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'

def add_logo(para, logo_b, width_in=1.8, sa=0):
    clear_runs(para)
    fmt_para(para, sb=0, sa=sa)
    run = para.add_run()
    run.add_picture(io.BytesIO(logo_b), width=Inches(width_in))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── Step 1: Extract logo bytes from para[1] ──────────────────────────────────────
ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
logo_bytes = None

for elem in doc.paragraphs[1]._element.iter():
    tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
    if tag == 'blip':
        rId = elem.get(f'{{{ns_r}}}embed')
        if rId and rId in doc.part.related_parts:
            logo_bytes = doc.part.related_parts[rId]._blob
            break

if not logo_bytes:
    print("ERROR: Could not extract logo image. Aborting.")
    sys.exit(1)

print(f"Logo extracted: {len(logo_bytes):,} bytes")

# ── Step 2: Rebuild Page 1 ───────────────────────────────────────────────────────
# Original indices (from scan):
# [0] empty
# [1] IMAGE (logo)
# [2] 'Islamic University of Lebanon'  ← redundant text, delete
# [3] 'Faculty of Engineering'          ← redundant text, delete
# [4] empty
# [5] empty                             ← extra, delete
# [6] 'OrionLead AI'
# [7] 'By'
# [8] 'Ali Jradeh'
# [9] 'Graduation Project Report'
# [10] 'Submitted...'
# [11] 'Bachelor of Engineering...'
# [12] 'Department of Computer Engineering'
# [13] 'Faculty of Engineering'
# [14] 'Supervised by'
# [15] 'Dr. Mohammad Alawwan'
# [16] empty                            ← delete
# [17] empty                            ← delete
# [18] 'Academic Year 2025-2026'
# [19] empty with sectPr               ← section boundary, do not touch content

p = doc.paragraphs

# Delete [0] (empty before logo)
delete_para(p[0]);  p = doc.paragraphs   # logo is now [0]

# Replace logo para [0] with properly sized logo + spacing
add_logo(p[0], logo_bytes, width_in=1.8, sa=16)

# Delete [1] 'Islamic University of Lebanon'
delete_para(p[1]);  p = doc.paragraphs

# Delete [1] 'Faculty of Engineering' (standalone redundant text)
delete_para(p[1]);  p = doc.paragraphs

# Delete one of the two remaining empty paras (now [1])
delete_para(p[1]);  p = doc.paragraphs

# Current layout after deletions:
# [0]  logo
# [1]  empty  ← keep as tiny spacer
# [2]  'OrionLead AI'
# [3]  'By'
# [4]  'Ali Jradeh'
# [5]  'Graduation Project Report'
# [6]  'Submitted in Partial Fulfillment...'
# [7]  'Bachelor of Engineering...'
# [8]  'Department of Computer Engineering'
# [9]  'Faculty of Engineering'
# [10] 'Supervised by'
# [11] 'Dr. Mohammad Alawwan'
# [12] empty
# [13] empty  ← delete
# [14] 'Academic Year 2025-2026'
# [15] section break para

fmt_para(p[1], sb=0, sa=0)                                                # [1] tiny spacer

set_text(p[2],  'OrionLead AI',                                            # [2]
         bold=True,  size=14, sb=8,  sa=6)

set_text(p[3],  'By',                                                      # [3]
         bold=False, size=12, sb=0,  sa=6)

set_text(p[4],  'Ali Jradeh',                                              # [4]
         bold=True,  size=12, sb=0,  sa=6)

set_text(p[5],  'Graduation Project Report',                               # [5]
         bold=True,  size=12, sb=18, sa=6)

set_text(p[6],  'Submitted in Partial Fulfillment of the Requirements for the Degree of',  # [6]
         bold=False, size=12, sb=24, sa=6)

set_text(p[7],  'Bachelor of Engineering in Computer Engineering',         # [7]
         bold=False, size=12, sb=0,  sa=18)

set_text(p[8],  'Department of Computer Engineering',                      # [8]
         bold=False, size=12, sb=0,  sa=6)

set_text(p[9],  'Faculty of Engineering',                                  # [9]
         bold=False, size=12, sb=0,  sa=18)

set_text(p[10], 'Supervised by',                                           # [10]
         bold=False, size=12, sb=0,  sa=6)

set_text(p[11], 'Dr. Mohammad Alawwan',                                    # [11]
         bold=False, size=12, sb=0,  sa=24)

# Delete two extra empty paragraphs [12] and [13]
delete_para(p[12]);  p = doc.paragraphs
delete_para(p[12]);  p = doc.paragraphs

set_text(p[12], 'Academic Year  2025 – 2026',                        # [12]
         bold=False, size=12, sb=0,  sa=0)

print("Page 1 rebuilt.")

# ── Step 3: Find section break paragraph ─────────────────────────────────────────
p = doc.paragraphs
section_idx = None
for i, para in enumerate(p):
    for elem in para._element.iter():
        if elem.tag == qn('w:sectPr'):
            section_idx = i
            break
    if section_idx is not None:
        break

if section_idx is None:
    print("ERROR: Could not find section break. Aborting page 2 rebuild.")
    sys.exit(1)

print(f"Section break at para [{section_idx}]: '{p[section_idx].text[:40]}'")
print(f"  Next para [{section_idx+1}]: '{p[section_idx+1].text[:40]}'")
print(f"  Next para [{section_idx+2}]: '{p[section_idx+2].text[:40]}'")

# ── Step 4: Rebuild Page 2 ───────────────────────────────────────────────────────
# After section break (section_idx):
# [s+1] empty             → logo
# [s+2] 'The Report Defense Committee...'
# [s+3] 'that this is the approved version...'
# [s+4] 'OrionLead AI'
# [s+5] 'Approved By:'
# [s+6] 'Supervisor Signature: ____'
# [s+7] empty             → delete
# [s+8] 'Examiner Signature: ____'
# [s+9] empty             → delete
# [s+10] 'Examiner Signature: ____'

p  = doc.paragraphs
s1 = section_idx + 1   # first empty on page 2 → logo
s2 = section_idx + 2   # 'The Report Defense Committee...'

# Logo at top of page 2 with large space_after to push body text to middle
add_logo(p[s1], logo_bytes, width_in=1.8, sa=120)

set_text(p[s2],   'The Report Defense Committee for Ali Jradeh Certifies',
         bold=True,  size=12, sb=0,  sa=6)

set_text(p[s2+1], 'that this is the approved version of the following report',
         bold=True,  size=12, sb=0,  sa=24)

set_text(p[s2+2], 'OrionLead AI',
         bold=True,  size=14, sb=0,  sa=48)

set_text(p[s2+3], 'Approved By:',
         bold=False, size=12, sb=0,  sa=18)

set_text(p[s2+4], 'Supervisor Signature:  ________________________________________',
         bold=False, size=12, sb=0,  sa=18)

# Delete empty gap between signatures
delete_para(p[s2+5]);  p = doc.paragraphs

set_text(p[s2+5], 'Examiner Signature:  ________________________________________',
         bold=False, size=12, sb=0,  sa=18)

delete_para(p[s2+6]);  p = doc.paragraphs

set_text(p[s2+6], 'Examiner Signature:  ________________________________________',
         bold=False, size=12, sb=0,  sa=0)

print("Page 2 rebuilt.")

# ── Save ─────────────────────────────────────────────────────────────────────────
doc.save(SRC)
print("\nDocument saved successfully!")
print(f"Backup is at: {BACKUP}")
