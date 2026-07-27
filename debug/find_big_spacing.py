"""
Find paragraphs with very large spacing or page-break-before that create blank-looking pages.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

doc = Document(r'C:\AI-Lead-Collection-System\report\OrionLead_AI_Report.docx')

print("Paragraphs with space_before > 60pt or space_after > 60pt:")
print("=" * 70)
for i, para in enumerate(doc.paragraphs):
    pf = para.paragraph_format
    sb = pf.space_before
    sa = pf.space_after
    pb_before = para._element.find(f'.//{qn("w:pageBreakBefore")}')

    # Convert to pt (1pt = 12700 EMU)
    sb_pt = round(sb.pt, 1) if sb else 0
    sa_pt = round(sa.pt, 1) if sa else 0

    has_break = pb_before is not None
    text = para.text.strip()[:60]

    if sb_pt > 60 or sa_pt > 60 or has_break:
        print(f"  [{i:3d}] sb={sb_pt:6.1f}pt sa={sa_pt:6.1f}pt pbBefore={has_break} | '{text}'")

print("\n\nEmpty paragraphs with non-zero spacing:")
print("=" * 70)
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        continue
    pf = para.paragraph_format
    sb = pf.space_before
    sa = pf.space_after
    sb_pt = round(sb.pt, 1) if sb else 0
    sa_pt = round(sa.pt, 1) if sa else 0
    if sb_pt > 10 or sa_pt > 10:
        print(f"  [{i:3d}] sb={sb_pt:6.1f}pt sa={sa_pt:6.1f}pt | (empty)")

print("\nDone.")
