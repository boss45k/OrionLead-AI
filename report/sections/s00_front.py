"""
s00_front.py — Front matter: Cover, Committee, Acknowledgments,
               Abstract, Table of Contents, List of Tables, List of Figures.
               Cover and committee pages match the CashPilot/IUL template with logo.
"""
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
import theme as T

LOGO = os.path.join(os.path.dirname(__file__), '..', 'iul_logo.png')


def build(doc):
    _cover(doc)
    _new_section(doc)     # section break → header/footer kicks in from here
    _committee(doc)
    _acknowledgments(doc)
    _abstract(doc)
    _toc(doc)
    _list_of_tables(doc)
    _list_of_figures(doc)


# ── Cover page — matches IUL CashPilot template ────────────────────────────────
def _cover(doc):
    # ── Logo centered at top ──
    T.spacer(doc, 8)
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_before = Pt(0)
    logo_p.paragraph_format.space_after  = Pt(10)
    if os.path.exists(LOGO):
        run = logo_p.add_run()
        run.add_picture(LOGO, width=Cm(5.5))

    # ── University name block ──
    T.cover_line(doc, "Islamic University of Lebanon",
                 size=16, bold=True, color=T.NAVY, before=2, after=2)
    T.cover_line(doc, "Faculty of Engineering",
                 size=13, bold=False, color=T.NAVY, before=0, after=18)

    # ── Thick separator ──
    T.h_rule(doc, T.NAVY_HEX, thick=16)
    T.spacer(doc, 22)

    # ── Project title (matches "CashPilot" style — centered, bold) ──
    T.cover_line(doc, "OrionLead AI",
                 size=26, bold=True, color=T.DARK, before=0, after=10)

    # ── "By" line ──
    T.cover_line(doc, "By",
                 size=13, bold=False, color=T.DARK, before=0, after=6)

    # ── Student name ──
    T.cover_line(doc, "Ali Jradeh",
                 size=16, bold=True, color=T.DARK, before=0, after=22)

    # ── Report label ──
    T.cover_line(doc, "Graduation Project Report",
                 size=14, bold=True, color=T.DARK, before=0, after=28)

    # ── Submission statement ──
    T.cover_line(doc,
                 "Submitted in Partial Fulfillment of the Requirements for the Degree of",
                 size=12, bold=True, color=T.DARK, before=0, after=4)
    T.cover_line(doc,
                 "Bachelor of Engineering in Computer Engineering",
                 size=12, bold=True, color=T.DARK, before=0, after=22)

    # ── Department / Faculty ──
    T.cover_line(doc, "Department of Computer Engineering",
                 size=13, bold=True, color=T.DARK, before=0, after=4)
    T.cover_line(doc, "Faculty of Engineering",
                 size=13, bold=True, color=T.DARK, before=0, after=20)

    # ── Supervised by ──
    T.cover_line(doc, "Supervised by",
                 size=12, bold=True, color=T.DARK, before=0, after=4)
    T.cover_line(doc, "Dr. Mohammad Alawwan",
                 size=13, bold=True, color=T.DARK, before=0, after=14)

    # ── Bottom thin rule + year ──
    T.h_rule(doc, T.NAVY_HEX, thick=8)
    T.spacer(doc, 6)
    T.cover_line(doc, "Academic Year  2025 – 2026",
                 size=12, bold=False, color=T.MID, before=0, after=0)


# ── Section break (cover → body with header/footer) ───────────────────────────
def _new_section(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sectPr = OxmlElement('w:sectPr')
    pgSz   = OxmlElement('w:pgSz')
    pgSz.set(qn('w:w'), '11906')
    pgSz.set(qn('w:h'), '16838')
    sectPr.append(pgSz)
    pgMar = OxmlElement('w:pgMar')
    pgMar.set(qn('w:top'),    '1418')
    pgMar.set(qn('w:right'),  '1134')
    pgMar.set(qn('w:bottom'), '1418')
    pgMar.set(qn('w:left'),   '1701')
    pgMar.set(qn('w:header'), '709')
    pgMar.set(qn('w:footer'), '709')
    sectPr.append(pgMar)
    pPr.append(sectPr)


# ── Defense Committee — matches IUL template ───────────────────────────────────
def _committee(doc):
    doc.add_page_break()

    # Logo at top (same as cover)
    T.spacer(doc, 8)
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_before = Pt(0)
    logo_p.paragraph_format.space_after  = Pt(14)
    if os.path.exists(LOGO):
        run = logo_p.add_run()
        run.add_picture(LOGO, width=Cm(4.5))

    T.spacer(doc, 10)

    # Approval statement (italic centered — matches template)
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after  = Pt(6)
    T.fr(p1.add_run("The Report Defense Committee for Ali Jradeh Certifies"),
         T.HEAD_FONT, 12, bold=True, italic=True, color=T.DARK)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(24)
    T.fr(p2.add_run("that this is the approved version of the following report"),
         T.HEAD_FONT, 12, bold=True, italic=True, color=T.DARK)

    # Project title
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after  = Pt(40)
    T.fr(p3.add_run("OrionLead AI"),
         T.HEAD_FONT, 18, bold=True, color=T.DARK)

    # "Approved By:" label
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(0)
    p4.paragraph_format.space_after  = Pt(24)
    T.fr(p4.add_run("Approved By:"),
         T.HEAD_FONT, 12, bold=True, color=T.DARK)

    # Signature lines (matching template format)
    _sig_line(doc, "Supervisor Signature:")
    T.spacer(doc, 14)
    _sig_line(doc, "Examiner Signature:")
    T.spacer(doc, 14)
    _sig_line(doc, "Examiner Signature:")


def _sig_line(doc, label):
    """Signature line: 'Label: ___________' left-aligned with indent."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.left_indent  = Cm(3.0)

    T.fr(p.add_run(f"{label}  "), T.HEAD_FONT, 12, bold=True, color=T.DARK)
    T.fr(p.add_run("_" * 40),     T.HEAD_FONT, 12, bold=False, color=T.DARK)


# ── Acknowledgments ────────────────────────────────────────────────────────────
def _acknowledgments(doc):
    doc.add_page_break()
    _front_chapter_heading(doc, "ACKNOWLEDGMENTS")

    T.body(doc,
        "First and foremost, all praise and gratitude are due to Allah, the Almighty, "
        "for granting me the strength, patience, wisdom, and perseverance required to "
        "complete this work. Without His grace, none of this would have been possible.")

    T.body(doc,
        "I would like to express my deepest and most sincere gratitude to my supervisor, "
        "DR. Mohammad Alawwan, whose expert guidance, unwavering support, and invaluable "
        "technical insight were instrumental throughout every phase of this project. His "
        "dedication to academic excellence, patience in addressing my questions, and "
        "consistent encouragement pushed me to deliver the highest quality of work. It has "
        "been a true privilege to work under his supervision.")

    T.body(doc,
        "I am also sincerely grateful to the Faculty of Engineering at the Islamic "
        "University of Lebanon for providing the academic foundation, resources, and "
        "opportunities that made this project possible. The knowledge, skills, and critical "
        "thinking abilities developed throughout my studies here form the intellectual "
        "backbone of everything built in this system.")

    T.body(doc,
        "To my family — my parents, siblings, and loved ones — thank you for your "
        "unconditional love, endless patience, and unwavering support throughout this "
        "journey. Every late night spent coding, every moment of doubt, and every "
        "challenge overcome was made bearable by the warmth of your encouragement and "
        "belief in me. I dedicate this achievement to you.")

    T.body(doc,
        "To my friends and colleagues who offered technical advice, participated in "
        "testing the application, provided honest feedback, and offered moral support "
        "during the most demanding periods of this project — I am truly and deeply "
        "grateful. Your enthusiasm and belief in OrionLead AI motivated me to push "
        "through every obstacle.")

    T.body(doc,
        "Finally, I wish to acknowledge the open-source community whose collective "
        "contributions to tools and frameworks including Flask, React, React Native, "
        "XGBoost, and Expo made it possible for a single developer to build and deliver "
        "a system of this scale and sophistication.")

    T.spacer(doc, 20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    T.fr(p.add_run("Ali Jradeh\nIslamic University of Lebanon\n2025 – 2026"),
         T.HEAD_FONT, 11, italic=True, color=T.MID)


# ── Abstract ───────────────────────────────────────────────────────────────────
def _abstract(doc):
    doc.add_page_break()
    _front_chapter_heading(doc, "ABSTRACT")

    T.body(doc,
        "The rapid evolution of artificial intelligence has opened transformative "
        "possibilities for business development and sales automation. OrionLead AI is a "
        "full-stack, AI-powered B2B (Business-to-Business) lead generation and "
        "qualification platform designed to automate the entire lifecycle of identifying, "
        "collecting, evaluating, and managing potential business leads. Traditional lead "
        "generation approaches are labor-intensive, inconsistent, and heavily dependent "
        "on manual effort, resulting in lost opportunities and reduced sales productivity.")

    T.body(doc,
        "OrionLead AI addresses these challenges through an intelligent, multi-layer "
        "pipeline that combines rule-based filtering, machine learning, and large language "
        "models (LLMs) to deliver high-quality, pre-scored leads to sales professionals. "
        "The system is composed of three tightly integrated components: a Python Flask "
        "RESTful backend powered by a four-layer AI qualification cascade (Anti-Junk Rules "
        "-> XGBoost ML -> Google Gemini LLM -> Groq Fallback), a React 18 web dashboard "
        "with real-time analytics, and a React Native cross-platform mobile application "
        "with offline-first synchronization via Supabase.")

    T.body(doc,
        "The AI qualification engine uses a trained XGBoost model (xgb_v5) with 37 "
        "engineered features derived from lead contact data, firmographic signals, "
        "behavioral indicators, and source metadata. Leads are scored on a 0-100 scale "
        "and classified into four tiers: Hot (>=80), Warm (>=60), Cold (>=30), and "
        "Unqualified (<30). An anti-junk engine with 12 hard-reject rules prevents fake, "
        "spam, or low-quality contacts from entering the pipeline. The platform integrates "
        "with eight external APIs including Hunter.io, Apollo.io, People Data Labs, "
        "Serper.dev, Google Places, ZeroBounce, Google Gemini, and Groq.")

    T.body(doc,
        "A Company-First Intelligence Pipeline evaluates organizational signals — website "
        "quality, business classification, and growth indicators — before resolving "
        "individual contacts, dramatically improving lead relevance. A QuickLabelBar "
        "enables one-click outcome labeling (Converted / Replied / No Reply / Not a Fit) "
        "that feeds a continuous ML retraining loop. The system features role-based access "
        "control (Admin, Manager, User), real-time push notifications via Firebase Cloud "
        "Messaging, CSV export, dark/light mode theming, and an automated CI/CD pipeline "
        "with 563 pytest tests validated against both SQLite and MySQL in parallel jobs "
        "on GitHub Actions.")

    T.spacer(doc, 10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    T.fr(p.add_run("Keywords: "), T.HEAD_FONT, 11, bold=True, color=T.NAVY)
    T.fr(p.add_run(
        "B2B Lead Generation · Artificial Intelligence · Machine Learning · XGBoost · "
        "Large Language Models · Flask · React Native · Lead Qualification · "
        "Sales Automation · Multi-Layer AI Pipeline"),
        T.BODY_FONT, 11, italic=True, color=T.MID)


# ── Table of Contents ──────────────────────────────────────────────────────────
def _toc(doc):
    doc.add_page_break()
    _front_chapter_heading(doc, "TABLE OF CONTENTS")

    entries = [
        ("Defense Committee Approval",                      "ii",   0, False),
        ("Acknowledgments",                                 "iii",  0, False),
        ("Abstract",                                        "iv",   0, False),
        ("Table of Contents",                               "v",    0, False),
        ("List of Tables",                                  "vi",   0, False),
        ("List of Figures",                                 "vii",  0, False),
        ("CHAPTER 1 — INTRODUCTION",                        "1",    0, True),
        ("1.1   Background and Motivation",                 "1",    1, False),
        ("1.2   Problem Statement",                         "3",    1, False),
        ("1.3   Proposed Solution",                         "5",    1, False),
        ("1.4   Project Objectives",                        "6",    1, False),
        ("1.5   Project Scope",                             "7",    1, False),
        ("1.6   Report Organization",                       "8",    1, False),
        ("CHAPTER 2 — LITERATURE REVIEW",                   "9",    0, True),
        ("2.1   Overview of B2B Lead Generation",           "9",    1, False),
        ("2.2   Traditional Lead Generation Methods",       "11",   1, False),
        ("2.3   AI and ML in Lead Qualification",           "13",   1, False),
        ("2.4   Existing Systems and Tools",                "16",   1, False),
        ("2.5   Research Gap and Motivation",               "18",   1, False),
        ("CHAPTER 3 — SYSTEM ANALYSIS",                     "19",   0, True),
        ("3.1   System External Entities",                  "19",   1, False),
        ("3.2   User Roles and Permissions",                "21",   1, False),
        ("3.3   Functional Requirements",                   "22",   1, False),
        ("3.4   Non-Functional Requirements",               "24",   1, False),
        ("3.5   Data Flow Diagram (DFD)",                   "25",   1, False),
        ("3.6   Use Case Diagram",                          "27",   1, False),
        ("3.7   Use Case Descriptions",                     "28",   1, False),
        ("CHAPTER 4 — SYSTEM DESIGN AND ARCHITECTURE",      "33",   0, True),
        ("4.1   Overall System Architecture",               "33",   1, False),
        ("4.2   Backend Architecture",                      "35",   1, False),
        ("4.3   AI Qualification Engine",                   "37",   1, False),
        ("4.4   Web Frontend Architecture",                 "40",   1, False),
        ("4.5   Mobile Application Architecture",           "42",   1, False),
        ("4.6   Database Design",                           "44",   1, False),
        ("4.7   Entity-Relationship Diagram (ERD)",         "50",   1, False),
        ("4.8   Security Architecture",                     "51",   1, False),
        ("4.9   External Integrations",                     "53",   1, False),
        ("CHAPTER 5 — SYSTEM IMPLEMENTATION",               "55",   0, True),
        ("5.1   Development Environment",                   "55",   1, False),
        ("5.2   Backend Implementation",                    "57",   1, False),
        ("5.3   AI Engine Implementation",                  "60",   1, False),
        ("5.4   Web Frontend Implementation",               "63",   1, False),
        ("5.5   Mobile Application Implementation",         "65",   1, False),
        ("5.6   Database Implementation",                   "67",   1, False),
        ("5.7   API Endpoints",                             "68",   1, False),
        ("5.8   Real-Time Synchronization",                 "69",   1, False),
        ("CHAPTER 6 — TESTING AND EVALUATION",              "71",   0, True),
        ("6.1   Testing Strategy",                          "71",   1, False),
        ("6.2   Functional Test Cases",                     "72",   1, False),
        ("6.3   API Testing",                               "76",   1, False),
        ("6.4   Non-Functional Testing",                    "78",   1, False),
        ("6.5   CI/CD Pipeline",                            "79",   1, False),
        ("CHAPTER 7 — SYSTEM RESULTS AND USER INTERFACE",   "81",   0, True),
        ("7.1   Web Dashboard — Authentication",            "81",   1, False),
        ("7.2   Web Dashboard — Lead Management",           "82",   1, False),
        ("7.3   Web Dashboard — Analytics",                 "83",   1, False),
        ("7.4   Web Dashboard — AI Collection",             "84",   1, False),
        ("7.5   Mobile Application — Leads",                "85",   1, False),
        ("7.6   Mobile Application — Analytics",            "86",   1, False),
        ("7.7   Mobile Application — Dark Mode",            "87",   1, False),
        ("CHAPTER 8 — CONCLUSION AND FUTURE WORK",          "88",   0, True),
        ("8.1   Conclusion",                                "88",   1, False),
        ("8.2   Achievements",                              "89",   1, False),
        ("8.3   Limitations",                               "91",   1, False),
        ("8.4   Future Work",                               "92",   1, False),
        ("REFERENCES",                                      "95",   0, True),
    ]
    for text, page, level, bold in entries:
        T.toc_entry(doc, text, page, level, bold)


# ── List of Tables ─────────────────────────────────────────────────────────────
def _list_of_tables(doc):
    doc.add_page_break()
    _front_chapter_heading(doc, "LIST OF TABLES")

    tables = [
        ("Table 2.1",  "Comparison of OrionLead AI with Existing Systems",              "17"),
        ("Table 3.1",  "System External Entities",                                      "19"),
        ("Table 3.2",  "User Roles and Permissions Matrix",                             "21"),
        ("Table 3.3",  "Functional Requirements",                                       "22"),
        ("Table 3.4",  "Non-Functional Requirements",                                   "24"),
        ("Table 3.5",  "Use Case: UC-01 Register Account",                              "28"),
        ("Table 3.6",  "Use Case: UC-02 Collect Leads with AI",                         "29"),
        ("Table 3.7",  "Use Case: UC-03 Qualify Lead with AI",                          "30"),
        ("Table 3.8",  "Use Case: UC-04 Label Lead Outcome",                            "31"),
        ("Table 3.9",  "Use Case: UC-05 Export Leads to CSV",                           "32"),
        ("Table 4.1",  "API Blueprints and Route Prefixes",                             "35"),
        ("Table 4.2",  "Users Table Attributes",                                        "44"),
        ("Table 4.3",  "Leads Table Attributes",                                        "45"),
        ("Table 4.4",  "LeadActivities Table Attributes",                               "46"),
        ("Table 4.5",  "LeadOutcomes Table Attributes",                                 "47"),
        ("Table 4.6",  "DataSources Table Attributes",                                  "47"),
        ("Table 4.7",  "ClassificationCategories Table Attributes",                     "48"),
        ("Table 4.8",  "SeenContacts Table Attributes",                                 "48"),
        ("Table 4.9",  "SyncLogs Table Attributes",                                     "49"),
        ("Table 4.10", "External API Integrations",                                     "53"),
        ("Table 5.1",  "Backend Technology Stack",                                      "55"),
        ("Table 5.2",  "Web Frontend Technology Stack",                                 "56"),
        ("Table 5.3",  "Mobile Technology Stack",                                       "56"),
        ("Table 5.4",  "XGBoost Feature Categories (37 features)",                      "61"),
        ("Table 5.5",  "Anti-Junk Rule Descriptions",                                   "62"),
        ("Table 5.6",  "Lead Management API Endpoints",                                 "68"),
        ("Table 6.1",  "Functional Test Cases — Authentication",                        "72"),
        ("Table 6.2",  "Functional Test Cases — Lead Management",                       "73"),
        ("Table 6.3",  "Functional Test Cases — AI Qualification",                      "75"),
        ("Table 6.4",  "API Endpoint Test Results",                                     "76"),
        ("Table 6.5",  "Non-Functional Test Results",                                   "78"),
    ]
    for num, title, page in tables:
        T.toc_entry(doc, f"{num}  —  {title}", page, level=0, bold=False)


# ── List of Figures ────────────────────────────────────────────────────────────
def _list_of_figures(doc):
    doc.add_page_break()
    _front_chapter_heading(doc, "LIST OF FIGURES")

    figures = [
        ("Figure 3.1",  "Level-0 Data Flow Diagram (Context Diagram)",          "25"),
        ("Figure 3.2",  "Level-1 Data Flow Diagram",                            "26"),
        ("Figure 3.3",  "Use Case Diagram",                                     "27"),
        ("Figure 4.1",  "Overall Three-Tier System Architecture",               "33"),
        ("Figure 4.2",  "Four-Layer AI Qualification Cascade",                  "37"),
        ("Figure 4.3a", "Core Domain Entity-Relationship Diagram (ERD)",        "50"),
        ("Figure 4.3b", "Supporting Tables Entity-Relationship Diagram",        "51"),
        ("Figure 4.4",  "Company-First Intelligence Pipeline",                  "52"),
        ("Figure 4.5",  "Real-Time Synchronization Architecture",               "54"),
        ("Figure 5.1",  "Backend Directory Structure",                          "57"),
        ("Figure 5.2",  "Mobile Application Directory Structure",               "65"),
        ("Figure 7.1",  "Web Dashboard — Login Screen",                         "81"),
        ("Figure 7.2",  "Web Dashboard — Leads List View",                      "82"),
        ("Figure 7.3",  "Web Dashboard — Lead Detail and QuickLabelBar",        "82"),
        ("Figure 7.4",  "Web Dashboard — Analytics Overview Dashboard",         "83"),
        ("Figure 7.5",  "Web Dashboard — AI Collection Panel",                  "84"),
        ("Figure 7.6",  "Mobile App — Leads Screen (Light Mode)",               "85"),
        ("Figure 7.7",  "Mobile App — Lead Detail and QuickLabelBar",           "85"),
        ("Figure 7.8",  "Mobile App — Analytics Screen",                        "86"),
        ("Figure 7.9",  "Mobile App — Dark Mode (All Screens)",                 "87"),
    ]
    for num, title, page in figures:
        T.toc_entry(doc, f"{num}  —  {title}", page, level=0, bold=False)


# ── Shared: front-matter chapter heading ───────────────────────────────────────
def _front_chapter_heading(doc, title):
    T.h_rule(doc, T.NAVY_HEX, 20)
    T.spacer(doc, 8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    T.fr(p.add_run(title), T.HEAD_FONT, 16, bold=True, color=T.NAVY)
    T.spacer(doc, 4)
    T.h_rule(doc, T.GOLD_HEX, 6)
    T.spacer(doc, 14)
