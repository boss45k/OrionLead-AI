"""
s09_references.py — References
Target: 2-3 pages
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
import theme as T
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build(doc):
    doc.add_page_break()
    T.h_rule(doc, T.NAVY_HEX, 20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    T.fr(p.add_run("REFERENCES"), T.HEAD_FONT, 16, bold=True, color=T.NAVY)

    T.h_rule(doc, T.GOLD_HEX, 6)
    T.spacer(doc, 12)

    refs = [
        # ── Lead Generation & B2B Sales ──────────────────────────────────────
        ("[1]  Kotler, P., & Keller, K. L. (2016). Marketing Management (15th ed.). Pearson "
         "Education. — Foundational framework for B2B lead qualification and pipeline management."),

        ("[2]  Lilien, G. L., & Rangaswamy, A. (2004). Marketing Engineering: Computer-Assisted "
         "Marketing Analysis and Planning (2nd ed.). Trafford Publishing. — Quantitative models "
         "for prospect scoring and sales funnel optimization."),

        ("[3]  Hunter.io. (2024). Email Finder API Documentation. Retrieved from "
         "https://hunter.io/api-documentation — Technical reference for the domain search "
         "and email verification endpoints used in the collection engine."),

        ("[4]  Apollo.io. (2024). Apollo API Reference (v1). Retrieved from "
         "https://apolloio.github.io/apollo-api-docs — Contact and company search API "
         "integrated as a primary data source."),

        ("[5]  People Data Labs. (2024). Person Enrichment API. Retrieved from "
         "https://docs.peopledatalabs.com — Enrichment API providing seniority, skills, "
         "and employment history fields used as XGBoost features."),

        # ── Machine Learning & XGBoost ────────────────────────────────────────
        ("[6]  Chen, T., & Guestrin, C. (2016). XGBoost: A Scalable Tree Boosting System. "
         "In Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge "
         "Discovery and Data Mining (pp. 785–794). ACM. "
         "https://doi.org/10.1145/2939672.2939785 — Original paper for the gradient "
         "boosting library used in the ML qualification layer."),

        ("[7]  Lundberg, S. M., & Lee, S.-I. (2017). A Unified Approach to Interpreting "
         "Model Predictions. In Advances in Neural Information Processing Systems, 30 "
         "(pp. 4765–4774). — SHAP framework referenced in Section 8.4.5 for future "
         "explainability work."),

        ("[8]  Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. "
         "Journal of Machine Learning Research, 12, 2825–2830. — Used for preprocessing "
         "pipelines, cross-validation, and evaluation metrics during model training."),

        ("[9]  Géron, A. (2022). Hands-On Machine Learning with Scikit-Learn, Keras, and "
         "TensorFlow (3rd ed.). O'Reilly Media. — Practical reference for feature engineering "
         "and model evaluation methodology applied in Chapter 5."),

        # ── Large Language Models & NLP ──────────────────────────────────────
        ("[10] Google. (2024). Gemini API Documentation. Google AI for Developers. Retrieved "
         "from https://ai.google.dev/docs — API reference for the Gemini 1.5 Flash model "
         "used in the LLM qualification layer and query parsing."),

        ("[11] Brown, T. B., et al. (2020). Language Models are Few-Shot Learners. "
         "In Advances in Neural Information Processing Systems, 33 (pp. 1877–1901). — "
         "Foundational work on prompt-based LLM usage referenced in the cascade design."),

        ("[12] Wei, J., et al. (2022). Chain-of-Thought Prompting Elicits Reasoning in "
         "Large Language Models. In Advances in Neural Information Processing Systems, 35. — "
         "Structured reasoning approach applied in the Gemini qualification prompt design."),

        # ── Web & Mobile Development ──────────────────────────────────────────
        ("[13] Facebook Inc. (2024). React Documentation (v18). Retrieved from "
         "https://react.dev — Official reference for the React 18 framework used in "
         "the web dashboard."),

        ("[14] Expo. (2024). Expo SDK 54 Documentation. Retrieved from "
         "https://docs.expo.dev — Official reference for the Expo managed workflow "
         "used in the mobile application."),

        ("[15] Ant Design. (2024). Ant Design 5 Component Library. Retrieved from "
         "https://ant.design/components/overview — UI component library reference for "
         "the web dashboard interface."),

        ("[16] Supabase. (2024). Supabase Realtime Documentation. Retrieved from "
         "https://supabase.com/docs/guides/realtime — Reference for the WebSocket "
         "subscription bridge used for mobile real-time synchronization."),

        # ── Backend & Infrastructure ──────────────────────────────────────────
        ("[17] Grinberg, M. (2018). Flask Web Development: Developing Web Applications "
         "with Python (2nd ed.). O'Reilly Media. — Reference for Flask application "
         "structure, blueprints, and request lifecycle used in the API layer."),

        ("[18] Celery Project. (2024). Celery: Distributed Task Queue (v5.3). Retrieved "
         "from https://docs.celeryq.dev — Reference for the asynchronous task execution "
         "used in collection jobs and background ML retraining."),

        ("[19] Google. (2024). Firebase Cloud Messaging (FCM) Documentation. Retrieved "
         "from https://firebase.google.com/docs/cloud-messaging — Reference for the "
         "push notification delivery system integrated into the mobile application."),

        # ── Software Engineering & Security ──────────────────────────────────
        ("[20] OWASP Foundation. (2021). OWASP Top Ten 2021. Retrieved from "
         "https://owasp.org/Top10 — Security standard applied during API layer design "
         "and penetration testing phase described in Chapter 6."),

        ("[21] Fowler, M. (2018). Refactoring: Improving the Design of Existing Code "
         "(2nd ed.). Addison-Wesley. — Software design principles applied in the "
         "modular architecture of the collection engine and qualification cascade."),

        ("[22] Kim, G., Humble, J., Debois, P., & Willis, J. (2016). The DevOps Handbook. "
         "IT Revolution Press. — CI/CD pipeline design principles reflected in the "
         "GitHub Actions dual-job (SQLite + MySQL) testing strategy."),
    ]

    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Pt(28)
        p.paragraph_format.first_line_indent = Pt(-28)
        from docx.shared import Pt as _Pt
        run = p.add_run(ref)
        run.font.name = "Times New Roman"
        run.font.size = _Pt(11)
